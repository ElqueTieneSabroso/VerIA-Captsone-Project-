from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parents[1]
SOURCE = WORKSPACE / "DOSSIER_CALIDAD_Y_CONTEXTO_VERIA.md"
OUTPUT = WORKSPACE / "VERIA_Dossier_Tecnico_y_Calidad_2026-07-24.docx"
LOGO = WORKSPACE / "assets" / "LOGO_2.png"
DATABASE_DIAGRAM = Path(r"C:\Users\IRowen\Downloads\Database V1.png")

PRESET_NAME = "compact_reference_guide"
HEADER_PATTERN = "editorial_cover"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203748"
MUTED = "66788A"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "E7F3EA"
PALE_YELLOW = "FFF4CE"
PALE_RED = "FCE8E6"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in CELL_MARGINS_DXA.items():
        tag = qn(f"w:{edge}")
        element = tc_mar.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def create_numbering_instance(document: Document, style_name="List Number") -> int:
    style = document.styles[style_name]._element
    style_num_pr = style.pPr.numPr
    base_num_id = int(style_num_pr.numId.val)
    numbering = document.part.numbering_part.element
    base_num = numbering.xpath(f'./w:num[@w:numId="{base_num_id}"]')[0]
    abstract_num_id = int(base_num.abstractNumId.val)
    existing_ids = [int(element.get(qn("w:numId"))) for element in numbering.num_lst]
    new_num_id = max(existing_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return new_num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_element])
    p_pr.append(num_pr)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {TABLE_WIDTH_DXA}: {widths_dxa}")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")


def calculate_widths(rows: list[list[str]]) -> list[int]:
    column_count = max(len(row) for row in rows)
    lengths = []
    for column in range(column_count):
        values = [row[column] if column < len(row) else "" for row in rows]
        lengths.append(max(7, min(42, max(len(value) for value in values))))

    minimum = 780 if column_count >= 5 else 1000
    available = TABLE_WIDTH_DXA - minimum * column_count
    total_weight = sum(lengths)
    widths = [
        minimum + round(available * weight / total_weight) for weight in lengths
    ]
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    return widths


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_border_bottom(paragraph, color=BLUE, size=8, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_font(run, size=9, color=MUTED)


def configure_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.5), WD_TAB_ALIGNMENT.RIGHT
    )
    left = paragraph.add_run("VERIA")
    set_font(left, size=9, color=MUTED, bold=True)
    right = paragraph.add_run("\tDossier técnico y evidencia de calidad")
    set_font(right, size=9, color=MUTED)
    paragraph_border_bottom(paragraph, color=PALE_BLUE, size=8, space=4)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = footer_paragraph.add_run("Página ")
    set_font(prefix, size=9, color=MUTED)
    add_page_field(footer_paragraph)

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    first_footer.paragraphs[0].text = ""


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.15

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(8.5)
    code_style.paragraph_format.left_indent = Inches(0.16)
    code_style.paragraph_format.right_indent = Inches(0.16)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(6)
    code_style.paragraph_format.line_spacing = 1.05


def add_inline_runs(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|https?://[^\s)]+)")
    for token in pattern.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Consolas", size=9.5, color=DARK_BLUE)
        else:
            run = paragraph.add_run(token)
            set_font(run)


def add_body_paragraph(document: Document, text: str, style=None):
    paragraph = document.add_paragraph(style=style)
    add_inline_runs(paragraph, text)
    return paragraph


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F4F6F9")
    p_pr.append(shading)
    run = paragraph.add_run("\n".join(lines))
    set_font(run, name="Consolas", size=8.5, color=INK)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (column_count - len(row)) for row in rows]
    widths = calculate_widths(normalized_rows)
    table = document.add_table(rows=len(normalized_rows), cols=column_count)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])

    for row_index, row_values in enumerate(normalized_rows):
        for column_index, value in enumerate(row_values):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(1.5)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline_runs(paragraph, value)
            for run in paragraph.runs:
                set_font(
                    run,
                    size=8.5,
                    bold=row_index == 0,
                    color=INK,
                )
            if row_index == 0:
                set_cell_shading(cell, PALE_BLUE)
            elif re.search(r"\b(Aprobada|Aprobado|Corregida)\b", value, re.I):
                set_cell_shading(cell, PALE_GREEN)
            elif re.search(r"\b(Pendiente|abierta|avisos)\b", value, re.I):
                set_cell_shading(cell, PALE_YELLOW)
            elif re.search(r"\b(Fallida|fallido)\b", value, re.I):
                set_cell_shading(cell, PALE_RED)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(42)

    if LOGO.exists():
        logo_paragraph = document.add_paragraph()
        logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = logo_paragraph.add_run()
        run.add_picture(str(LOGO), width=Inches(1.45))
        for doc_pr in run._r.xpath(".//wp:docPr"):
            doc_pr.set("descr", "Logotipo de VERIA")
            doc_pr.set("title", "VERIA")
        logo_paragraph.paragraph_format.space_after = Pt(18)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("DOSSIER TÉCNICO · CORTE 24 JUL 2026")
    set_font(run, size=10, color=BLUE, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("VERIA")
    set_font(run, size=30, color=INK, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(5)
    run = subtitle.add_run("Dossier técnico, registro de cambios")
    set_font(run, size=15, color=DARK_BLUE, bold=True)

    subtitle_two = document.add_paragraph()
    subtitle_two.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_two.paragraph_format.space_after = Pt(24)
    run = subtitle_two.add_run("y evidencia de calidad")
    set_font(run, size=15, color=DARK_BLUE, bold=True)

    purpose = document.add_paragraph()
    purpose.alignment = WD_ALIGN_PARAGRAPH.CENTER
    purpose.paragraph_format.left_indent = Inches(0.7)
    purpose.paragraph_format.right_indent = Inches(0.7)
    purpose.paragraph_format.space_after = Pt(30)
    run = purpose.add_run(
        "Fuente de verdad para continuidad técnica y para la redacción "
        "posterior de un documento escalable de aproximadamente 200 páginas."
    )
    set_font(run, size=10.5, color=MUTED, italic=True)

    metrics = document.add_table(rows=2, cols=3)
    values = [
        ("68", "pruebas aprobadas"),
        ("21/21", "comprobaciones Expo"),
        ("4", "flujos Maestro preparados"),
    ]
    for column, (value, label) in enumerate(values):
        top = metrics.cell(0, column)
        bottom = metrics.cell(1, column)
        top.text = value
        bottom.text = label
        for cell in (top, bottom):
            set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(0)
        for run in top.paragraphs[0].runs:
            set_font(run, size=18, color=BLUE, bold=True)
        for run in bottom.paragraphs[0].runs:
            set_font(run, size=8.5, color=MUTED, bold=True)
    set_repeat_table_header(metrics.rows[0])
    set_table_geometry(metrics, [3120, 3120, 3120])

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(26)
    meta.paragraph_format.space_after = Pt(0)
    run = meta.add_run(
        f"Preset editorial: {PRESET_NAME}  ·  Cabecera: {HEADER_PATTERN}"
    )
    set_font(run, size=8.5, color=MUTED)
    document.add_page_break()


def add_database_model_figure(document: Document) -> None:
    if not DATABASE_DIAGRAM.exists():
        paragraph = document.add_paragraph()
        add_inline_runs(
            paragraph,
            "[EVIDENCIA PENDIENTE] No se encontró la imagen del modelo conceptual "
            "de datos al regenerar el documento.",
        )
        return

    figure = document.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_before = Pt(4)
    figure.paragraph_format.space_after = Pt(4)
    figure.paragraph_format.keep_with_next = True
    run = figure.add_run()
    run.add_picture(str(DATABASE_DIAGRAM), width=Inches(6.35))
    for doc_pr in run._r.xpath(".//wp:docPr"):
        doc_pr.set(
            "descr",
            "Diagrama conceptual de VERIA con Usuarios como entidad central, "
            "preferencias de accesibilidad, auditivas, hápticas e interacción, "
            "además de relaciones con Imágenes y Errores.",
        )
        doc_pr.set("title", "Modelo conceptual de datos de VERIA")

    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True
    caption_run = caption.add_run(
        "Figura 1. Modelo conceptual aportado por el equipo. Se evalúa como "
        "diseño histórico; no equivale por sí solo al esquema físico conectado."
    )
    set_font(caption_run, size=8.5, color=MUTED, italic=True)


def add_contents(document: Document, section_titles: list[str]) -> None:
    heading = document.add_paragraph("Contenido y mapa de lectura", style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    intro = document.add_paragraph()
    add_inline_runs(
        intro,
        "El documento distingue resultados ejecutados, validaciones estáticas, "
        "pendientes e incidencias abiertas. Las secciones 22 y 23 contienen el "
        "plan editorial y el prompt transferible; las secciones 27 a 35 "
        "concilian el corpus documental incorporado.",
    )

    toc_num_id = create_numbering_instance(document)
    for title in section_titles:
        paragraph = document.add_paragraph(style="List Number")
        apply_numbering(paragraph, toc_num_id)
        clean_title = re.sub(r"^\d+\.\s*", "", title)
        add_inline_runs(paragraph, clean_title)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), PALE_YELLOW)
    p_pr.append(shading)
    run = paragraph.add_run(
        "Lectura recomendada: resumen ejecutivo → incidencias → resultados → "
        "trazabilidad → corpus documental → paquete de contexto para otro chat."
    )
    set_font(run, size=9.5, color=INK, bold=True)
    document.add_page_break()


def parse_markdown(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    start_index = next(
        index for index, line in enumerate(lines) if line.startswith("## 1.")
    )
    lines = lines[start_index:]
    paragraph_buffer: list[str] = []
    in_code = False
    code_lines: list[str] = []
    first_major = True
    active_numbering_id = None
    active_list_paragraph = None

    def flush_paragraph():
        if paragraph_buffer:
            text = " ".join(part.strip() for part in paragraph_buffer).strip()
            if text:
                add_body_paragraph(document, text)
            paragraph_buffer.clear()

    index = 0
    while index < len(lines):
        line = lines[index]

        if line.startswith("```"):
            flush_paragraph()
            active_numbering_id = None
            active_list_paragraph = None
            if in_code:
                add_code_block(document, code_lines)
                code_lines.clear()
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if line.strip() == "[[FIGURE:database-conceptual-model]]":
            flush_paragraph()
            active_numbering_id = None
            active_list_paragraph = None
            add_database_model_figure(document)
            index += 1
            continue

        if line.startswith("|"):
            flush_paragraph()
            active_numbering_id = None
            active_list_paragraph = None
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            table_rows = [
                parse_table_row(table_line)
                for table_line in table_lines
                if not is_table_separator(table_line)
            ]
            add_markdown_table(document, table_rows)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading_match:
            flush_paragraph()
            active_numbering_id = None
            active_list_paragraph = None
            marks, title = heading_match.groups()
            page_break_before = False
            if len(marks) == 2:
                if not first_major:
                    page_break_before = True
                first_major = False
                style = "Heading 1"
            elif len(marks) == 3:
                style = "Heading 2"
            else:
                style = "Heading 3"
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.page_break_before = page_break_before
            add_inline_runs(paragraph, title)
            index += 1
            continue

        bullet_match = re.match(r"^\s*-\s+(.+)$", line)
        numbered_match = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if bullet_match or numbered_match:
            flush_paragraph()
            style = "List Bullet" if bullet_match else "List Number"
            text = (bullet_match or numbered_match).group(1)
            paragraph = document.add_paragraph(style=style)
            if numbered_match:
                if active_numbering_id is None:
                    active_numbering_id = create_numbering_instance(document)
                apply_numbering(paragraph, active_numbering_id)
            else:
                active_numbering_id = None
            add_inline_runs(paragraph, text)
            active_list_paragraph = paragraph
            index += 1
            continue

        if line.strip() in {"---", ""}:
            flush_paragraph()
            active_numbering_id = None
            active_list_paragraph = None
            index += 1
            continue

        if active_list_paragraph is not None and line[:1].isspace():
            add_inline_runs(active_list_paragraph, f" {line.strip()}")
            index += 1
            continue

        active_numbering_id = None
        active_list_paragraph = None
        paragraph_buffer.append(line)
        index += 1

    flush_paragraph()
    if in_code and code_lines:
        add_code_block(document, code_lines)


def add_document_end(document: Document) -> None:
    heading = document.add_paragraph("Ficha de cierre", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.space_before = Pt(0)
    rows = [
        ["Campo", "Valor"],
        ["Documento fuente", SOURCE.name],
        ["Documento generado", OUTPUT.name],
        ["Preset", PRESET_NAME],
        ["Patrón de portada", HEADER_PATTERN],
        ["Fecha de corte", "24 de julio de 2026"],
        ["Pruebas aprobadas", "68"],
        ["E2E sobre dispositivo", "Pendiente; no se atribuye ejecución"],
        ["Incidencia principal", "Fidelidad semántica del modelo visual"],
    ]
    add_markdown_table(document, rows)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(
        "Fin del dossier · conservar junto con la fuente Markdown y la evidencia."
    )
    set_font(run, size=9, color=MUTED, italic=True)


def enable_update_fields(document: Document) -> None:
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def audit_document(document: Document) -> dict:
    heading_sequence = []
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            try:
                heading_sequence.append(int(paragraph.style.name.split()[-1]))
            except ValueError:
                pass
    skipped = []
    previous = 0
    for index, level in enumerate(heading_sequence):
        if previous and level > previous + 1:
            skipped.append({"index": index, "from": previous, "to": level})
        previous = level

    tables = []
    for index, table in enumerate(document.tables):
        widths = []
        first_row = table.rows[0]
        for cell in first_row.cells:
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            widths.append(int(tc_w.get(qn("w:w"))) if tc_w is not None else 0)
        tables.append(
            {
                "index": index,
                "columns": len(first_row.cells),
                "width_sum": sum(widths),
            }
        )

    return {
        "preset": PRESET_NAME,
        "header_pattern": HEADER_PATTERN,
        "headings": len(heading_sequence),
        "heading_skips": skipped,
        "tables": len(document.tables),
        "table_geometry_failures": [
            item for item in tables if item["width_sum"] != TABLE_WIDTH_DXA
        ],
        "paragraphs": len(document.paragraphs),
    }


def build() -> Path:
    markdown = SOURCE.read_text(encoding="utf-8")
    section_titles = re.findall(r"^##\s+(\d+\..+)$", markdown, flags=re.MULTILINE)

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    configure_styles(document)
    configure_header_footer(section)
    enable_update_fields(document)

    properties = document.core_properties
    properties.title = "VERIA — Dossier técnico y evidencia de calidad"
    properties.subject = (
        "Arquitectura, cambios, pruebas, incidencias y contexto transferible"
    )
    properties.author = "Proyecto VERIA"
    properties.keywords = (
        "VERIA, Expo, React Native, Whisper, Ollama, Maestro, pruebas, accesibilidad"
    )
    properties.comments = (
        "Generado desde DOSSIER_CALIDAD_Y_CONTEXTO_VERIA.md con el preset "
        f"{PRESET_NAME} y la portada {HEADER_PATTERN}."
    )

    add_cover(document)
    add_contents(document, section_titles)
    parse_markdown(document, markdown)
    add_document_end(document)

    report = audit_document(document)
    if report["heading_skips"]:
        raise RuntimeError(f"Heading hierarchy failures: {report['heading_skips']}")
    if report["table_geometry_failures"]:
        raise RuntimeError(
            f"Table geometry failures: {report['table_geometry_failures']}"
        )

    document.save(OUTPUT)
    print(report)
    print(f"Created: {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    try:
        build()
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise
