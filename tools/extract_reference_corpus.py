from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


SOURCE_FILES = [
    "Informes técnicos.pptx",
    "Pruebas Automatizadas.pptx",
    "Tipos de pruebas manuales.pptx",
    "Pruebas de software(1).pptx",
    "Metodologías y documentación en pruebas de software.pptx",
    "Software Development Life Cycle.pdf",
    "PERT - PERT.pdf",
    "ALVARADO-DE_AVILA-RIVERA(1).docx",
    "BASE DE DATOS VERIA.docx",
    "ALVARADO-DE_AVILA-RIVERA.docx",
    "Avances 17_Julio_2026_.docx",
    "DE AVILA - RIVERA .docx",
    "DE AVILA - RIVERA(1).docx",
    "De_Avila-Rivera 3_B.docx",
    "IA integration(1).docx",
    "IA integration.docx",
    "Interview_Guide - Jorge.docx",
    "Organigrama _ Calendario.xlsx",
    "Software Design and Analysis .pdf",
    "Software Design and Analysis .docx",
    "manuel1.pdf",
    "Database V1.png",
]

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def natural_key(name: str) -> list[Any]:
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", name)
    ]


def pptx_text(path: Path) -> dict[str, Any]:
    slides: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        slide_names = sorted(
            (
                name
                for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            ),
            key=natural_key,
        )
        note_names = {
            int(match.group(1)): name
            for name in archive.namelist()
            if (
                match := re.fullmatch(
                    r"ppt/notesSlides/notesSlide(\d+)\.xml", name
                )
            )
        }
        media = [
            name
            for name in archive.namelist()
            if name.startswith("ppt/media/") and not name.endswith("/")
        ]
        for index, slide_name in enumerate(slide_names, start=1):
            root = ET.fromstring(archive.read(slide_name))
            text_runs = [
                (node.text or "").strip()
                for node in root.findall(".//a:t", NS)
                if (node.text or "").strip()
            ]
            note_runs: list[str] = []
            note_name = note_names.get(index)
            if note_name:
                note_root = ET.fromstring(archive.read(note_name))
                note_runs = [
                    (node.text or "").strip()
                    for node in note_root.findall(".//a:t", NS)
                    if (node.text or "").strip()
                ]
            slides.append(
                {
                    "slide": index,
                    "text": text_runs,
                    "notes": note_runs,
                }
            )
    return {
        "slide_count": len(slides),
        "media_count": len(media),
        "slides": slides,
    }


def iter_docx_blocks(document: Document) -> list[dict[str, Any]]:
    body = document.element.body
    paragraphs = {paragraph._p: paragraph for paragraph in document.paragraphs}
    tables = {table._tbl: table for table in document.tables}
    blocks: list[dict[str, Any]] = []
    for child in body.iterchildren():
        if child in paragraphs:
            paragraph = paragraphs[child]
            text = paragraph.text.strip()
            if text:
                blocks.append(
                    {
                        "kind": "paragraph",
                        "style": paragraph.style.name if paragraph.style else "",
                        "text": text,
                    }
                )
        elif child in tables:
            table = tables[child]
            rows = []
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            if any(any(cell for cell in row) for row in rows):
                blocks.append({"kind": "table", "rows": rows})
    return blocks


def docx_text(path: Path) -> dict[str, Any]:
    document = Document(path)
    blocks = iter_docx_blocks(document)
    headings = [
        block
        for block in blocks
        if block["kind"] == "paragraph"
        and str(block.get("style", "")).lower().startswith("heading")
    ]
    with zipfile.ZipFile(path) as archive:
        image_count = len(
            [
                name
                for name in archive.namelist()
                if name.startswith("word/media/") and not name.endswith("/")
            ]
        )
        pages = None
        if "docProps/app.xml" in archive.namelist():
            app_root = ET.fromstring(archive.read("docProps/app.xml"))
            for node in app_root.iter():
                if node.tag.endswith("Pages") and node.text:
                    try:
                        pages = int(node.text)
                    except ValueError:
                        pages = None
    return {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "heading_count": len(headings),
        "image_count": image_count,
        "stored_page_count": pages,
        "blocks": blocks,
    }


def pdf_text(path: Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    pages: list[dict[str, Any]] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append(
            {
                "page": index,
                "text": text,
                "characters": len(text),
            }
        )
    metadata = {
        str(key): str(value)
        for key, value in (reader.metadata or {}).items()
        if value is not None
    }
    return {
        "page_count": len(reader.pages),
        "metadata": metadata,
        "pages": pages,
        "text_pages": sum(1 for page in pages if page["characters"] > 20),
    }


def xlsx_text(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        values: list[dict[str, Any]] = []
        formula_count = 0
        nonempty = 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                nonempty += 1
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                if len(values) < 5000:
                    values.append(
                        {
                            "cell": cell.coordinate,
                            "value": json_value(cell.value),
                        }
                    )
        sheets.append(
            {
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "nonempty_cells": nonempty,
                "formula_count": formula_count,
                "values": values,
                "truncated": nonempty > len(values),
            }
        )
    workbook.close()
    return {"sheet_count": len(sheets), "sheets": sheets}


def image_info(path: Path) -> dict[str, Any]:
    from PIL import Image

    with Image.open(path) as image:
        return {
            "width": image.width,
            "height": image.height,
            "format": image.format,
            "mode": image.mode,
        }


def main() -> None:
    downloads = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    output.mkdir(parents=True, exist_ok=True)
    corpus = []
    for file_name in SOURCE_FILES:
        path = downloads / file_name
        item: dict[str, Any] = {
            "name": file_name,
            "path": str(path),
            "exists": path.exists(),
        }
        if not path.exists():
            corpus.append(item)
            continue
        item.update(
            {
                "extension": path.suffix.lower(),
                "bytes": path.stat().st_size,
                "modified": datetime.fromtimestamp(
                    path.stat().st_mtime
                ).isoformat(),
                "sha256": sha256(path),
            }
        )
        try:
            if path.suffix.lower() == ".pptx":
                item["content"] = pptx_text(path)
            elif path.suffix.lower() == ".docx":
                item["content"] = docx_text(path)
            elif path.suffix.lower() == ".pdf":
                item["content"] = pdf_text(path)
            elif path.suffix.lower() == ".xlsx":
                item["content"] = xlsx_text(path)
            elif path.suffix.lower() in {".png", ".jpg", ".jpeg"}:
                item["content"] = image_info(path)
        except Exception as error:  # noqa: BLE001
            item["error"] = f"{type(error).__name__}: {error}"
        corpus.append(item)

    output_path = output / "reference-corpus.json"
    output_path.write_text(
        json.dumps(corpus, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    summary = []
    for item in corpus:
        content = item.get("content", {})
        summary.append(
            {
                "name": item["name"],
                "exists": item["exists"],
                "error": item.get("error"),
                "pages": content.get("page_count")
                or content.get("stored_page_count"),
                "slides": content.get("slide_count"),
                "sheets": content.get("sheet_count"),
                "paragraphs": content.get("paragraph_count"),
                "tables": content.get("table_count"),
                "images": content.get("image_count"),
            }
        )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"Corpus: {output_path}")


if __name__ == "__main__":
    main()
